from fastapi import FastAPI, File, UploadFile, HTTPException, Form, BackgroundTasks
from fastapi.responses import StreamingResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
import os
import io
import json
import tempfile
import re
import base64
from typing import Optional
import traceback
from pathlib import Path
from dotenv import load_dotenv
import httpx

# Document processing
from docx import Document
import google.generativeai as genai
from PIL import Image
try:
    import fitz  # PyMuPDF for PDF handling
except ImportError:
    fitz = None

# Supabase
from supabase import create_client, Client

# Load environment variables
load_dotenv()

app = FastAPI(title="Document Processing API")

# Add CORS middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # For development
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Configure Gemini API
try:
    genai.configure(api_key=os.environ.get("GEMINI_API_KEY"))
    GEMINI_MODEL = "gemini-2.0-flash"  # Using older model name for better compatibility
    GEMINI = genai.GenerativeModel(GEMINI_MODEL)
    print(f"Gemini model {GEMINI_MODEL} configured successfully")
except Exception as e:
    print(f"Error configuring Gemini: {e}")
    GEMINI = None

# Configure Supabase client
try:
    supabase: Client = create_client(
        os.environ.get("SUPABASE_URL", ""),
        os.environ.get("SUPABASE_KEY", "")
    )
    print("Supabase client configured successfully")
except Exception as e:
    print(f"Error configuring Supabase: {e}")
    supabase = None

# Regular expression for matching placeholders
PLACEHOLDER_RE = re.compile(r"\{\{(.*?)\}\}")

# Helper functions
def convert_pdf_to_images(pdf_bytes):
    """Convert PDF pages to a list of image bytes."""
    if not fitz:
        raise Exception("PyMuPDF not installed. Install with: pip install PyMuPDF")
    
    images = []
    try:
        pdf_document = fitz.open(stream=pdf_bytes, filetype="pdf")
        for page_num in range(len(pdf_document)):
            page = pdf_document.load_page(page_num)
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))  # 2x zoom for better quality
            img_bytes = pix.tobytes("png")
            images.append(img_bytes)
    except Exception as e:
        raise Exception(f"Failed to convert PDF to images: {e}")
    return images

def convert_docx_to_text(docx_bytes):
    """Extract text from a DOCX file."""
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp.write(docx_bytes)
        tmp.flush()
        doc = Document(tmp.name)
        text = "\n".join([para.text for para in doc.paragraphs])
        return text

def extract_json_from_text(text):
    """Extract JSON from text using regex pattern matching."""
    # Try to find JSON pattern in the text
    json_match = re.search(r'(\{.*\})', text, re.DOTALL)
    if json_match:
        try:
            return json.loads(json_match.group(1))
        except json.JSONDecodeError:
            pass
    return None

@app.get("/")
def read_root():
    """Health check endpoint"""
    return {"status": "ok", "message": "Document Processing API is running"}

@app.post("/api/extract")
async def extract(
    file: UploadFile = File(...),
    template_id: Optional[str] = Form(None)
):
    """
    Extract information from uploaded documents using OCR.
    
    Args:
        file: The document file (PDF, DOCX, or image)
        template_id: Optional template ID to use (defaults to hardcoded value if not provided)
    
    Returns:
        JSON containing extracted values and missing fields
    """
    if not GEMINI:
        raise HTTPException(500, "Gemini API not configured")
    
    if not supabase:
        raise HTTPException(500, "Supabase connection not available")
    
    # Use provided template_id or fallback to hardcoded value
    template_id = template_id or "9fc0c5fc-2885-4d58-ba0f-4711244eb7df"
    
    try:
        # 1. Get template from Supabase
        try:
            tpl_row = supabase.table("templates").select("*")\
                    .eq("id", template_id)\
                    .single().execute().data
            
            if not tpl_row:
                raise HTTPException(404, f"Template not found with ID: {template_id}")
            
            placeholder_json: dict = tpl_row["info_json"]["required_fields"]
            print(f"Found template {template_id} with {len(placeholder_json)} placeholders")
        except Exception as db_err:
            print(f"Database error: {db_err}")
            raise HTTPException(500, f"Error fetching template: {str(db_err)}")
        
        # 2. Read the uploaded file
        file_content = await file.read()
        file_extension = Path(file.filename).suffix.lower()
        print(f"Processing file: {file.filename} ({file_extension})")
        
        # 3. Process the file based on its type
        image_bytes = None
        
        if file_extension in (".jpg", ".jpeg", ".png"):
            image_bytes = file_content
            print(f"Detected image file: {len(image_bytes)} bytes")
        elif file_extension == ".pdf" and fitz:
            # For demo purposes, we'll just use the first page
            try:
                image_bytes = convert_pdf_to_images(file_content)[0]
                print(f"Converted first PDF page to image: {len(image_bytes)} bytes")
            except Exception as pdf_err:
                print(f"PDF conversion error: {pdf_err}")
                raise HTTPException(400, f"Failed to process PDF: {str(pdf_err)}")
        elif file_extension in (".doc", ".docx"):
            # For DOCX, we'll just extract text for now
            text = convert_docx_to_text(file_content)
            return JSONResponse(
                status_code=400,
                content={
                    "error": "DOCX processing not fully implemented yet",
                    "extracted_text": text[:1000] + "..." if len(text) > 1000 else text
                }
            )
        else:
            raise HTTPException(400, f"Unsupported file type: {file_extension}")
        
        if not image_bytes:
            raise HTTPException(400, "Could not process file")
        
        # 4. Ask Gemini to extract information
        # Prepare the keys from the placeholder JSON
        all_keys = [k.strip('{}') for k in placeholder_json.keys()]
        
        # Create a dictionary mapping keys to their descriptions for better context
        field_descriptions = {k.strip('{}'): placeholder_json[k] for k in placeholder_json.keys()}
        
        # First prompt: Extract raw data from the image
        initial_prompt = f"""You are an extraction engine for Philippine PSA birth certificates.
Extract information from the provided image, focusing on the following fields:

{", ".join(all_keys)}

Return ONLY a JSON object with these exact keys. For example:
{{
    "first_name": "Juan",
    "last_name": "Dela Cruz",
    ...
}}

If a value is not found or unclear, use an empty string.
DO NOT include any explanations or additional text outside the JSON object.
"""
        
        try:
            print("Sending initial extraction request to Gemini API...")
            # Properly create and prepare the PIL image object
            img = Image.open(io.BytesIO(image_bytes))
            
            # Ensure the image is in a compatible format (RGB)
            if img.mode not in ['RGB', 'L']:
                img = img.convert('RGB')
            
            initial_response = GEMINI.generate_content(
                contents=[initial_prompt, img],
                generation_config={
                    "temperature": 0.1,
                    "max_output_tokens": 2048,
                }
            )
            
            # Get text response
            initial_response_text = initial_response.text
            print(f"Received initial response of length {len(initial_response_text)}")
            
            # Try to extract JSON from the response
            raw_extracted = extract_json_from_text(initial_response_text)
            if not raw_extracted:
                print("Failed to extract JSON from initial response")
                print(f"Response preview: {initial_response_text[:200]}...")
                # Return error with raw response for debugging
                return {
                    "template_id": template_id,
                    "error": "Could not extract valid JSON from Gemini initial response",
                    "raw_response": initial_response_text[:500] + "..." if len(initial_response_text) > 500 else initial_response_text,
                    "extracted_ocr": {},
                    "missing_value_keys": placeholder_json
                }
            
            # 5. Process the extracted data with a second LLM call to organize it correctly
            # Create a comprehensive description of the fields for better processing
            field_info = "\n".join([f"{key}: {field_descriptions[key]}" for key in all_keys])
            
            # Explain tricky fields that need special handling
            special_fields_info = """
Special fields to check:
1. Checkbox fields: These should contain "X" if checked, empty string if not checked:
   - m (male checkbox)
   - f (female checkbox)
   - typeb1, typeb2, typeb3 (birth type checkboxes)
   - mb_first, mb_second, mb_third (multiple birth order checkboxes)

2. Name fields: Ensure names are placed in the correct fields:
   - first_name, middle_name, last_name (for child)
   - mother_first_name, mother_middle_name, mother_last_name (for mother)
   - father_first_name, father_middle_name, father_last_name (for father)
   
3. Checkbox data should never contain names, addresses, or other text data
"""
            
            # Make refinement prompt
            refine_prompt = f"""You are an expert in organizing extracted data from Philippine PSA birth certificates.
            
I have extracted raw data from a birth certificate, but some fields might be incorrectly populated.
Please reorganize this data to ensure it's correctly placed in the appropriate fields.

Here's the description of each field:
{field_info}

{special_fields_info}

Here's the raw extracted data:
{json.dumps(raw_extracted, indent=2)}

- For gender checkboxes ("m" and "f"), only one should have "X" based on the gender, not names or other data
- Make sure dates are properly formatted
- Ensure names are in the correct fields (don't mix up child, mother, and father names)
- For all checkbox fields (like "typeb1", "m", "f", etc.), use "X" if checked, empty string if not
- If names were incorrectly extracted into checkbox fields, move them to the appropriate name fields

Return a corrected JSON object with the same keys but properly organized values.
"""
            
            # Make the second LLM call to refine the data
            refine_response = GEMINI.generate_content(
                contents=refine_prompt,
                generation_config={
                    "temperature": 0.1,
                    "max_output_tokens": 4096,
                }
            )
            
            refine_response_text = refine_response.text
            print(f"Received refinement response of length {len(refine_response_text)}")
            
            # Extract the refined JSON
            refined_json = extract_json_from_text(refine_response_text)
            if not refined_json:
                print("Failed to extract JSON from refinement response, falling back to raw extraction")
                refined_json = raw_extracted
            
            # 6. Format the response with placeholders
            extracted = {}
            for k in placeholder_json.keys():
                clean_key = k.strip("{}")
                if clean_key in refined_json and refined_json[clean_key]:
                    extracted[k] = refined_json[clean_key]
            
            missing = {k: placeholder_json[k] for k in placeholder_json if k not in extracted}
            
            return {
                "template_id": template_id,
                "doc_type": tpl_row.get("doc_type", ""),
                "variation": tpl_row.get("variation", ""),
                "extracted_ocr": extracted,
                "missing_value_keys": missing
            }
        
        except Exception as e:
            print(f"Gemini API error: {e}")
            traceback.print_exc()
            return {
                "template_id": template_id,
                "error": f"Error processing with Gemini: {str(e)}",
                "extracted_ocr": {},
                # "missing_value_keys": placeholder_json
            }
    
    except Exception as e:
        print(f"Extraction error: {e}")
        traceback.print_exc()
        raise HTTPException(500, f"Error during extraction: {str(e)}")

if __name__ == "__main__":
    import uvicorn
    # Load port from environment or default to 8000
    port = int(os.environ.get("PORT", 8000))
    # Run the FastAPI app
    uvicorn.run(app, host="0.0.0.0", port=port)