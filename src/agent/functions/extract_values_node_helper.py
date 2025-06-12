# backend/src/agent/functions/extract_values_node_helper.py

import os
import io
import json
import tempfile
import re
import base64
from typing import Optional, Dict, Any, Union, List, Set
import traceback
from pathlib import Path
from dotenv import load_dotenv
import httpx

# Document processing
from docx import Document
import google.generativeai as genai
from PIL import Image
try:
    import fitz  # PyMuPDF for PDF handling
except ImportError:
    fitz = None

# Supabase
from supabase import create_client, Client

# Load environment variables
load_dotenv()

class DocumentExtractor:
    def __init__(self):
        """Initialize the document extractor with API configurations."""
        # Configure Gemini API
        try:
            genai.configure(api_key=os.environ.get("GEMINI_API_KEY"))
            self.gemini_model = "gemini-2.0-flash"
            self.gemini = genai.GenerativeModel(self.gemini_model)
            print(f"Gemini model {self.gemini_model} configured successfully")
        except Exception as e:
            print(f"Error configuring Gemini: {e}")
            self.gemini = None

        # Configure Supabase client
        try:
            self.supabase: Client = create_client(
                os.environ.get("SUPABASE_URL", ""),
                os.environ.get("SUPABASE_KEY", "")
            )
            print("Supabase client configured successfully")
        except Exception as e:
            print(f"Error configuring Supabase: {e}")
            self.supabase = None

        # Regular expression for matching placeholders
        self.placeholder_re = re.compile(r"\{\{(.*?)\}\}")

        # Define radio button groups - only one from each group should be selected
        self.radio_groups = {
            "gender": ["{ma}", "{fe}"],  # Male/Female
            "attendant": ["{at1}", "{at2}", "{at3}", "{at4}", "{at5}"],  # Different attendant types
            "birth_type": ["{tb1}", "{tb2}", "{tb3}"],  # Birth type (single, twin, triplet)
            "multiple_birth": ["{imb1}", "{imb2}", "{imb3}"],  # Multiple birth order
        }

        # Field categorization for better extraction context
        self.field_categories = {
            "child_info": [
                "{first_name}", "{middle_name}", "{last_name}", "{ma}", "{fe}", 
                "{day}", "{month}", "{year}", "{wab}", "{bo}"
            ],
            "birth_location": [
                "{placebirth_host}", "{pb_city/mnplt}", "{pb_prov}"
            ],
            "birth_type": [
                "{tb1}", "{tb2}", "{tb3}", "{imb1}", "{imb2}", "{imb3}", "{imb_o}"
            ],
            "mother_info": [
                "{m_first_name}", "{m_middle_name}", "{m_last_name}", "{m_occupation}",
                "{m_religion}", "{m_citizenship}", "{mtb}", "{m_city/mnplt}", "{m_prov}"
            ],
            "father_info": [
                "{f_first_name}", "{f_middle_name}", "{f_last_name}", "{f_occupation}",
                "{f_religion}", "{f_citizenship}", "{ftb}", "{residence}"
            ],
            "attendant": [
                "{at1}", "{at2}", "{at3}", "{at4}", "{at5}"
            ],
            "statistics": [
                "{tncba}", "{ncslib}", "{ncbobnd}"
            ],
            "registry_info": [
                "{city/mnplt}", "{province}", "{registration_n}", "{receive_by_date}",
                "{receive_by_name}", "{receive_by_title}"
            ],
            "informant": [
                "{informant_name}", "{informant_relation}", "{informant_address}", "{informant_date}"
            ],
            "certification": [
                "{cob_t}", "{cob_date}", "{cob_title}", "{cob_print}", "{cob_address}",
                "{prepared_by_name}", "{prepared_by_position}", "{prepared_by_date}"
            ],
            "parents_marriage": [
                "{date_place_marriage_parents}"
            ]
        }

    def identify_radio_groups_in_template(self, placeholder_json: Dict[str, str]) -> Dict[str, List[str]]:
        """
        Identify which radio groups are present in the current template.
        Returns only the groups that have fields in the template.
        """
        template_keys = set(placeholder_json.keys())
        active_groups = {}
        
        for group_name, group_fields in self.radio_groups.items():
            present_fields = [field for field in group_fields if field in template_keys]
            if len(present_fields) > 1:  # Only consider it a group if multiple fields are present
                active_groups[group_name] = present_fields
                
        return active_groups

    def apply_radio_button_logic(self, extracted: Dict[str, Any], missing: Dict[str, Any], 
                                active_groups: Dict[str, List[str]]) -> tuple[Dict[str, Any], Dict[str, Any]]:
        """
        Apply radio button logic: if one field in a group is selected, remove others from missing.
        """
        updated_missing = missing.copy()
        
        for group_name, group_fields in active_groups.items():
            # Check if any field in this group was extracted
            selected_fields = [field for field in group_fields if field in extracted and extracted[field]["value"]]
            
            if selected_fields:
                # Remove all other fields in this group from missing
                for field in group_fields:
                    if field != selected_fields[0] and field in updated_missing:
                        del updated_missing[field]
                        print(f"Removed {field} from missing (radio group: {group_name}, selected: {selected_fields[0]})")
        
        return extracted, updated_missing

    async def download_file_from_url(self, file_url: str) -> bytes:
        """Download file content from a public URL."""
        try:
            async with httpx.AsyncClient() as client:
                response = await client.get(file_url)
                response.raise_for_status()
                return response.content
        except Exception as e:
            raise Exception(f"Failed to download file from URL {file_url}: {str(e)}")

    def convert_pdf_to_images(self, pdf_bytes: bytes) -> list:
        """Convert PDF pages to a list of image bytes."""
        if not fitz:
            raise Exception("PyMuPDF not installed. Install with: pip install PyMuPDF")
        
        images = []
        try:
            pdf_document = fitz.open(stream=pdf_bytes, filetype="pdf")
            for page_num in range(len(pdf_document)):
                page = pdf_document.load_page(page_num)
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))  # 2x zoom for better quality
                img_bytes = pix.tobytes("png")
                images.append(img_bytes)
        except Exception as e:
            raise Exception(f"Failed to convert PDF to images: {e}")
        return images

    def convert_docx_to_text(self, docx_bytes: bytes) -> str:
        """Extract text from a DOCX file."""
        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            tmp.write(docx_bytes)
            tmp.flush()
            doc = Document(tmp.name)
            text = "\n".join([para.text for para in doc.paragraphs])
            return text

    def extract_json_from_text(self, text: str) -> Optional[Dict]:
        """Extract JSON from text using regex pattern matching."""
        # Try to find JSON pattern in the text
        json_match = re.search(r'(\{.*\})', text, re.DOTALL)
        if json_match:
            try:
                return json.loads(json_match.group(1))
            except json.JSONDecodeError:
                pass
        return None

    def create_birth_certificate_extraction_prompt(self, placeholder_json: Dict[str, str], 
                                                  active_groups: Dict[str, List[str]]) -> str:
        """Create a specialized extraction prompt for Philippine birth certificates."""
        
        # Create group descriptions
        group_descriptions = []
        for group_name, group_fields in active_groups.items():
            group_desc = f"\n**{group_name.upper()} GROUP (select only ONE):**\n"
            for field in group_fields:
                clean_key = field.strip('{}')
                desc = placeholder_json.get(field, {}).get('description', '') if isinstance(placeholder_json.get(field), dict) else placeholder_json.get(field, '')
                group_desc += f"  - {field}: {desc}\n"
            group_descriptions.append(group_desc)
        
        # Create field descriptions by category
        field_sections = []
        for category, fields in self.field_categories.items():
            category_fields = [f for f in fields if f in placeholder_json]
            if category_fields:
                section = f"\n**{category.upper().replace('_', ' ')} FIELDS:**\n"
                for field in category_fields:
                    field_info = placeholder_json.get(field, {})
                    if isinstance(field_info, dict):
                        label = field_info.get('label', field)
                        description = field_info.get('description', '')
                    else:
                        label = field
                        description = field_info
                    section += f"  {field}: {label} - {description}\n"
                field_sections.append(section)
        
        # Create the main prompt
        prompt = f"""You are an expert OCR extraction engine specialized in Philippine PSA Birth Certificates.

Extract information from the provided birth certificate image, focusing on the following categorized fields:

{"".join(field_sections)}

**IMPORTANT RADIO BUTTON LOGIC:**
The following fields are grouped - only ONE field from each group should have a value:
{"".join(group_descriptions)}

**EXTRACTION RULES FOR BIRTH CERTIFICATES:**

1. **Radio Button Groups**: Only select ONE option per group based on what you see checked
2. **Checkbox Fields**: Use "X" for checked boxes, empty string for unchecked
3. **Text Fields**: Extract exact text as it appears on the document
4. **Name Fields**: 
   - Separate first, middle, and last names for child, mother, and father
   - Do not mix up names between different people
5. **Date Fields**: Extract dates as they appear (day, month, year separately if requested)
6. **Numeric Fields**: Extract numbers only (ages, weights, counts)
7. **Address Fields**: Extract complete addresses as written
8. **Registry Information**: Look for official stamps, registry numbers, dates

**SPECIFIC FIELD GUIDANCE:**
- Sex: Look for checked boxes next to Male/Female
- Birth Type: Look for Single/Twin/Triplet checkboxes
- Attendant: Look for Physician/Nurse/Midwife/Hilot/Others checkboxes
- Multiple Birth Order: If twins/triplets, look for birth order indicators
- Weight at Birth: Usually in grams
- Ages: Mother's and Father's age at time of birth
- Registry Info: Official numbers, dates, and certifying official information

**OUTPUT FORMAT:**
Return ONLY a JSON object with the field keys and their extracted values. Use empty strings for fields not found or not applicable.

Example:
{{
    "first_name": "Juan",
    "middle_name": "Santos",
    "last_name": "Dela Cruz",
    "ma": "X",
    "fe": "",
    "day": "15",
    "month": "March",
    "year": "2020",
    "tb1": "X",
    "tb2": "",
    "at1": "X",
    "m_first_name": "Maria",
    "f_first_name": "Jose"
}}

DO NOT include explanations, notes, or any text outside the JSON object.
"""
        return prompt

    def create_refinement_prompt(self, raw_extracted: Dict, placeholder_json: Dict[str, str], 
                               active_groups: Dict[str, List[str]]) -> str:
        """Create a refinement prompt to clean up and organize extracted data."""
        
        # Create field type mapping for validation
        checkbox_fields = []
        text_fields = []
        
        for field, info in placeholder_json.items():
            field_info = info if isinstance(info, dict) else {'description': info}
            description = field_info.get('description', '').lower()
            
            if 'checkbox' in description and 'mark with x' in description:
                checkbox_fields.append(field)
            else:
                text_fields.append(field)
        
        # Create validation rules
        validation_rules = f"""
**FIELD TYPE VALIDATION:**

Checkbox Fields (should contain only "X" or empty string):
{', '.join(checkbox_fields)}

Text Fields (should contain actual text content):
{', '.join(text_fields)}

**RADIO BUTTON GROUPS (only ONE per group should have "X"):**
"""
        
        for group_name, group_fields in active_groups.items():
            validation_rules += f"\n{group_name.upper()}: {', '.join(group_fields)}"
        
        prompt = f"""You are a data validation expert for Philippine birth certificate extraction.

Review and correct the following extracted data to ensure it follows proper formatting rules:

{validation_rules}

**CURRENT EXTRACTED DATA:**
{json.dumps(raw_extracted, indent=2)}

**CORRECTION RULES:**
1. **Checkbox Fields**: Must contain only "X" (if checked) or "" (if unchecked)
2. **Radio Groups**: Only ONE field per group should contain "X"
3. **Name Fields**: Should contain proper names, not checkmarks or irrelevant text
4. **Date/Number Fields**: Should contain appropriate date or numeric values
5. **Address Fields**: Should contain complete address information
6. **Text Fields**: Should not contain "X" unless it's part of actual text content

**COMMON ERRORS TO FIX:**
- Names appearing in checkbox fields → Move to appropriate name fields
- Multiple checkboxes selected in radio groups → Keep only the most appropriate one
- "X" appearing in text fields that should contain actual text
- Missing or misplaced information that was extracted to wrong fields

Return a corrected JSON object with the same structure but properly formatted values.
"""
        return prompt

    async def extract_values_from_document(
        self, 
        template_id: str,
        base_file_public_url: str
    ) -> Dict[str, Any]:
        """
        Extract information from a document using OCR based on template requirements.
        
        Args:
            template_id: Template ID to get required fields from database
            base_file_public_url: Public URL of the document to process
        
        Returns:
            Dictionary containing extracted values and missing fields
        """
        if not self.gemini:
            raise Exception("Gemini API not configured")
        
        if not self.supabase:
            raise Exception("Supabase connection not available")
        
        try:
            # 1. Get template from Supabase
            try:
                tpl_row = self.supabase.table("templates").select("*")\
                        .eq("id", template_id)\
                        .single().execute().data
                
                if not tpl_row:
                    raise Exception(f"Template not found with ID: {template_id}")
                
                placeholder_json: dict = tpl_row["info_json"]["required_fields"]
                print(f"Found template {template_id} with {len(placeholder_json)} placeholders")
            except Exception as db_err:
                print(f"Database error: {db_err}")
                raise Exception(f"Error fetching template: {str(db_err)}")
            
            # 2. Identify active radio groups
            active_groups = self.identify_radio_groups_in_template(placeholder_json)
            print(f"Identified {len(active_groups)} active radio groups: {list(active_groups.keys())}")
            
            # 3. Download the file from the public URL
            try:
                file_content = await self.download_file_from_url(base_file_public_url)
                print(f"Downloaded file from URL: {len(file_content)} bytes")
            except Exception as download_err:
                print(f"Download error: {download_err}")
                raise Exception(f"Failed to download file: {str(download_err)}")
            
            # 4. Determine file type from URL or content
            file_extension = Path(base_file_public_url.split('?')[0]).suffix.lower()
            if not file_extension:
                # Try to determine from content-type or file signature if no extension
                if file_content.startswith(b'%PDF'):
                    file_extension = '.pdf'
                elif file_content.startswith(b'\xff\xd8\xff'):
                    file_extension = '.jpg'
                elif file_content.startswith(b'\x89PNG'):
                    file_extension = '.png'
                else:
                    file_extension = '.pdf'  # Default assumption
            
            print(f"Processing file with extension: {file_extension}")
            
            # 5. Process the file based on its type
            image_bytes = None
            
            if file_extension in (".jpg", ".jpeg", ".png"):
                image_bytes = file_content
                print(f"Detected image file: {len(image_bytes)} bytes")
            elif file_extension == ".pdf" and fitz:
                # For birth certificates, we'll process the first page
                try:
                    image_bytes = self.convert_pdf_to_images(file_content)[0]
                    print(f"Converted first PDF page to image: {len(image_bytes)} bytes")
                except Exception as pdf_err:
                    print(f"PDF conversion error: {pdf_err}")
                    raise Exception(f"Failed to process PDF: {str(pdf_err)}")
            elif file_extension in (".doc", ".docx"):
                # For DOCX, we'll just extract text for now
                text = self.convert_docx_to_text(file_content)
                return {
                    "template_id": template_id,
                    "error": "DOCX processing not fully implemented yet",
                    "extracted_text": text[:1000] + "..." if len(text) > 1000 else text,
                    "extracted_ocr": {},
                    "missing_value_keys": placeholder_json
                }
            else:
                raise Exception(f"Unsupported file type: {file_extension}")
            
            if not image_bytes:
                raise Exception("Could not process file")
            
            # 6. Initial extraction with specialized birth certificate prompt
            try:
                print("Sending birth certificate extraction request to Gemini API...")
                
                # Properly create and prepare the PIL image object
                img = Image.open(io.BytesIO(image_bytes))
                
                # Ensure the image is in a compatible format (RGB)
                if img.mode not in ['RGB', 'L']:
                    img = img.convert('RGB')
                
                # Create specialized birth certificate prompt
                initial_prompt = self.create_birth_certificate_extraction_prompt(placeholder_json, active_groups)
                
                initial_response = self.gemini.generate_content(
                    contents=[initial_prompt, img],
                    generation_config={
                        "temperature": 0.1,
                        "max_output_tokens": 4096,
                    }
                )
                
                # Get text response
                initial_response_text = initial_response.text
                print(f"Received initial response of length {len(initial_response_text)}")
                
                # Try to extract JSON from the response
                raw_extracted = self.extract_json_from_text(initial_response_text)
                if not raw_extracted:
                    print("Failed to extract JSON from initial response")
                    print(f"Response preview: {initial_response_text[:200]}...")
                    return {
                        "template_id": template_id,
                        "error": "Could not extract valid JSON from Gemini initial response",
                        "raw_response": initial_response_text[:500] + "..." if len(initial_response_text) > 500 else initial_response_text,
                        "extracted_ocr": {},
                        "missing_value_keys": placeholder_json
                    }
                
                # 7. Refinement pass to clean up the data
                try:
                    refine_prompt = self.create_refinement_prompt(raw_extracted, placeholder_json, active_groups)
                    
                    refine_response = self.gemini.generate_content(
                        contents=refine_prompt,
                        generation_config={
                            "temperature": 0.1,
                            "max_output_tokens": 4096,
                        }
                    )
                    
                    refine_response_text = refine_response.text
                    print(f"Received refinement response of length {len(refine_response_text)}")
                    
                    # Extract the refined JSON
                    refined_json = self.extract_json_from_text(refine_response_text)
                    if not refined_json:
                        print("Failed to extract JSON from refinement response, using raw extraction")
                        refined_json = raw_extracted
                
                except Exception as refine_err:
                    print(f"Refinement error: {refine_err}, using raw extraction")
                    refined_json = raw_extracted
                
                # 8. Format the response with placeholders and apply radio button logic
                extracted = {}
                for k in placeholder_json.keys():
                    clean_key = k.strip("{}")
                    field_info = placeholder_json[k]
                    
                    # Get label from field info
                    if isinstance(field_info, dict):
                        label = field_info.get('label', k)
                    else:
                        label = k
                    
                    # Check if value was extracted
                    if clean_key in refined_json and refined_json[clean_key]:
                        extracted[k] = {
                            "label": label,
                            "value": refined_json[clean_key]
                        }
                
                # Format missing fields with label and empty value
                missing = {}
                for k in placeholder_json.keys():
                    if k not in extracted:
                        field_info = placeholder_json[k]
                        if isinstance(field_info, dict):
                            label = field_info.get('label', k)
                        else:
                            label = k
                        
                        missing[k] = {
                            "label": label,
                            "value": ""
                        }
                
                # Apply radio button logic to clean up missing fields
                extracted, missing = self.apply_radio_button_logic(extracted, missing, active_groups)
                
                print(f"Final extraction: {len(extracted)} extracted, {len(missing)} missing")
                
                return {
                    "template_id": template_id,
                    "doc_type": tpl_row.get("doc_type", ""),
                    "variation": tpl_row.get("variation", ""),
                    "extracted_ocr": extracted,
                    "missing_value_keys": missing,
                    "success": True
                }
            
            except Exception as e:
                print(f"Gemini API error: {e}")
                traceback.print_exc()
                return {
                    "template_id": template_id,
                    "error": f"Error processing with Gemini: {str(e)}",
                    "extracted_ocr": {},
                    "missing_value_keys": placeholder_json,
                    "success": False
                }
        
        except Exception as e:
            print(f"Extraction error: {e}")
            traceback.print_exc()
            raise Exception(f"Error during extraction: {str(e)}")


# Convenience function for direct usage in the orchestrator
async def extract_values_from_document(
    template_id: str,
    base_file_public_url: str
) -> Dict[str, Any]:
    """
    Convenience function to extract values from a document using template requirements.
    
    Args:
        template_id: Template ID to get required fields from database
        base_file_public_url: Public URL of the document to process
    
    Returns:
        Dictionary containing extracted values and missing fields
    
    Example usage:
        result = await extract_values_from_document(
            template_id="9fc0c5fc-2885-4d58-ba0f-4711244eb7df",
            base_file_public_url="https://example.com/document.pdf"
        )
    """
    extractor = DocumentExtractor()
    return await extractor.extract_values_from_document(
        template_id=template_id,
        base_file_public_url=base_file_public_url
    )